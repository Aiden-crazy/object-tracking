# -*- coding: utf-8 -*-
"""
report_part2.py —— 报告内容·第二部分
第3章 系统设计与实现 / 第4章 系统测试 / 第5章 系统总结 / 第6章 心得体会 / 参考文献 / 附录
作者：【姓名】  学号：【学号】  创建时间：2026-07
"""
from docx.shared import Pt

from docx_lib import add_body, add_code, add_fig, add_h1, add_h2, add_h3, add_table

ROOT = None


def F(name):
    return ROOT + "/demo/figures/" + name


def S(name):
    return ROOT + "/demo/screenshots/" + name


# ============================ 第3章 ============================
def build_ch3(doc):
    add_h1(doc, "3 单目标跟踪系统的设计与实现")
    # ---------------- 3.1 总体设计 ----------------
    add_h2(doc, "3.1 总体设计")
    add_h3(doc, "3.1.1 设计思路与总体架构")
    add_body(doc,
             "依据任务书对系统流程的要求，本系统按“表现层—业务服务层—视觉处理层—数据层”四层"
             "架构进行设计，总体架构如图3-1所示。")
    add_fig(doc, F("fig1_1_系统总体架构图.png"), "图3-1 系统总体架构图")
    add_body(doc,
             "各层职责如下：表现层包括微信小程序（普通用户）与 Web 管理端（管理员）；业务服务层"
             "为 Spring Boot Web 后端，承担用户认证、文件存储、视觉任务调度与 REST API 服务；"
             "视觉处理层为独立的 Python + OpenCV 服务，通过 FastAPI 对外提供首帧提取与单目标跟踪"
             "接口；数据层默认使用内嵌 H2 数据库（MySQL 兼容模式，可切换 MySQL 8.0）保存用户与"
             "任务信息，使用磁盘文件系统保存原始素材（视频/图片）、结果文件（结果视频/结果图片）与"
             "跟踪日志。视觉处理层独立部署，通过 HTTP 接口与后端解耦，便于后续替换或升级算法。")

    add_h3(doc, "3.1.2 技术选型")
    add_table(doc, "表3-1 关键技术选型及理由",
              ["层次", "技术", "选型理由"],
              [
                  ["用户端", "微信原生小程序", "任务书指定小程序端；无需安装、生态成熟，wx.uploadFile"
                   " 天然支持文件上传"],
                  ["管理端", "原生 HTML+CSS+JS 单页", "零构建、零依赖，随后端静态资源发布，答辩演示"
                   "环境适应性强"],
                  ["后端", "Spring Boot 3 + MyBatis + JWT + BCrypt", "主流企业级组合，注解式 SQL 开发"
                   "高效；JWT 无状态鉴权、BCrypt 密码安全"],
                  ["数据库", "内嵌 H2（MySQL 兼容模式）；可选 MySQL 8.0", "任务书指定数据库功能；"
                   "默认内嵌 H2 免安装、随程序自动建库，跨电脑演示无障碍，保留 MySQL 配置一键切换"],
                  ["视觉", "Python + OpenCV(CSRT) + FastAPI", "OpenCV 跟踪器免训练开箱即用；"
                   "重检测采用多尺度灰度模板匹配 + HSV 直方图验证；FastAPI 异步高性能，便于后端调用"],
                  ["协作", "Git + GitHub", "任务书要求使用版本控制工具进行团队协作"],
              ],
              col_widths=[1.8, 4.6, 7.6])
    add_body(doc,
             "视觉跟踪算法选型时对比了主流跟踪器：KCF 基于核相关滤波，速度极快但难以处理尺度"
             "变化与遮挡；CSRT 在通道与空间可靠性上进行了改进，精度更高且支持尺度估计；MOSSE 速度"
             "最快但精度有限。考虑任务书要求“目标被遮挡后能再次捕捉”，本系统以 CSRT 作为主跟踪器，"
             "并为遮挡丢失场景叠加了“多尺度灰度模板匹配 + HSV 直方图相关度验证”的重检测机制；"
             "目标特征则由首帧目标区域的灰度模板（形状/纹理）与 HSV 二维直方图（颜色）共同表达。")

    add_h3(doc, "3.1.3 系统功能结构")
    add_body(doc,
             "系统功能结构如图3-2所示，自上而下分为用户端功能、管理端功能与视觉处理功能三大模块，"
             "各模块相互独立、通过后端 REST 接口与任务表协同。")
    add_fig(doc, F("fig_tree_功能结构图.png"), "图3-2 系统功能结构图")

    # ---------------- 3.2 详细设计 ----------------
    add_h2(doc, "3.2 详细设计")
    add_h3(doc, "3.2.1 视觉跟踪模块设计")
    add_body(doc,
             "视觉跟踪模块是本系统的核心。根据“跟踪—丢失—再捕捉”的业务要求，模块被设计为"
             "“初始化、正常跟踪、遮挡搜索、跟踪失败结束”四种状态的状态机，如图3-3所示。")
    add_fig(doc, F("fig3_1_跟踪器状态机图.png"), "图3-3 单目标跟踪器状态机", width_cm=13.5)
    add_body(doc,
             "各状态职责如下：")
    add_body(doc,
             "（1）初始化（INIT）：读取素材首帧（视频取第 0 帧，图片素材本身即首帧），通过交互框选"
             "（本地演示端鼠标框选或客户端传入 bbox）确定目标框；未框选时自动确定目标——视频采用"
             "“中值背景差 + Otsu 阈值 + 最大连通域”的运动目标检测，图片素材没有帧间运动信息，"
             "回退为画面中央区域。随后截取目标区域灰度图作为参考模板，同时统计目标区域的 HSV 颜色"
             "直方图作为外观参考模型，然后初始化 CSRT 跟踪器。")
    add_body(doc,
             "（2）正常跟踪（TRACKING）：逐帧调用 CSRT 更新目标位置；每 5 帧进行一次跟踪质量"
             "校验：计算当前跟踪框区域与参考模型的直方图相关度；同时逐帧检查跟踪框是否越界。"
             "相关度连续 3 次校验低于阈值或跟踪框连续越界 3 帧，判定目标丢失，转入搜索状态；"
             "未达到确认帧数前继续输出跟踪框（防抖，避免单帧相似度抖动就整帧转入全图搜索）。")
    add_body(doc,
             "（3）遮挡搜索（SEARCHING）：在全图中以参考模板按 0.85/1.0/1.15 三种尺度进行模板"
             "匹配，将最高匹配位置作为候选目标，再用 HSV 直方图相关度对候选进行验证；候选连续 2 帧"
             "命中后认为目标找回，重新初始化 CSRT 并恢复跟踪，同时累计“恢复次数”。若直至视频结束"
             "仍未找回，则状态置为跟踪失败结束。")
    add_body(doc,
             "（4）跟踪失败结束（LOST_FINAL）：记录丢失状态并输出统计信息。")
    add_body(doc,
             "视觉处理的总体流程如图3-4所示。跟踪过程中，模块逐帧在画面上绘制跟踪框、状态标识"
             "（TRACKING/SEARCHING）与目标运动轨迹，最终输出带标注的结果视频（图片素材则输出标注"
             "了目标框的结果图片）、逐帧轨迹日志（JSON）以及包含总帧数、跟踪帧数、搜索帧数、丢失"
             "次数、恢复次数与处理帧率等指标的统计信息。")
    add_fig(doc, F("fig3_2_视觉处理流程图.png"), "图3-4 视觉处理流程图", width_cm=10.5)

    add_h3(doc, "3.2.2 用户管理与任务调度模块设计")
    add_body(doc,
             "用户管理模块完成注册、登录（JWT 签发）、修改密码等用户端功能与查询、新增、删除、"
             "重置密码等管理端功能，密码统一使用 BCrypt 加盐哈希存储，任何接口均不返回密码字段；"
             "管理端接口通过拦截器校验令牌中的 ADMIN 角色，实现权限隔离。")
    add_body(doc,
             "任务调度模块处理“上传—处理—回存”的异步链路。为支持“先在首帧画面上框选目标、再提交"
             "跟踪”的交互，上传接口提供两段式：小程序以 defer=true 上传素材，后端将文件落盘至 "
             "uploads/origin 目录并在 t_task 表登记任务（PENDING），调用视觉服务 prepare 接口提取"
             "首帧图片（落盘到 uploads/frames）并返回视频原始尺寸与自动识别目标框，前端据此显示画面"
             "供用户拖拽画框；用户确认后调用 /api/task/{id}/start 提交目标框，后台单线程线程池取出"
             "任务，置为 PROCESSING 后调用视觉服务 track_local 接口（传递本地文件绝对路径，避免大文件"
             "二次传输）；处理完成后将结果文件（视频 .mp4 或图片 .jpg，按素材类型自动选择）与日志"
             "复制到 uploads/results 目录，更新任务为 SUCCESS 并写入统计 JSON；任一步骤异常则置 "
             "FAILED 并记录 error_msg。上传任务处理的详细设计如表3-2所示（参照模板表3.1格式）。")
    add_table(doc, "表3-2 “视频/图片上传与跟踪处理”功能详细设计说明",
              ["项目", "内容"],
              [
                  ["功能名称", "视频/图片上传与跟踪处理"],
                  ["涉及数据表", "t_user、t_task（含磁盘文件系统）"],
                  ["功能表述", "接收用户上传的视频或图片（可选 bbox），登记任务并异步调度视觉服务"
                   "完成单目标跟踪/目标定位，回存结果供查询播放"],
                  ["输入项", "素材文件（multipart/file）、可选首帧目标框 bbox(x,y,w,h)；"
                   "defer 参数控制是否只提取首帧（框选用）"],
                  ["业务处理描述", "① 校验文件类型与 bbox 格式；② 文件落盘并 INSERT t_task 状态"
                   "PENDING；③ defer=true 时调用视觉服务提首帧并返回 frameUrl/尺寸/自动框，等待 "
                   "start；④ start 或 defer=false 时线程池异步置 PROCESSING 并调用视觉服务；"
                   "⑤ 成功后复制结果文件、UPDATE 状态 SUCCESS 与统计 JSON；失败置 FAILED 并记录原因"],
                  ["输出项（正确）", "返回任务 ID；前端轮询至 SUCCESS 后可播放结果视频/查看结果图片、"
                   "查看统计"],
                  ["输出项（错误）", "返回明确错误提示（格式不支持、bbox 非法、视觉服务不可用等）"],
                  ["界面要求", "微信小程序页面 + Web 管理端页面"],
              ],
              col_widths=[3.2, 10.8])

    add_h3(doc, "3.2.3 数据库设计")
    add_body(doc,
             "数据库概念结构采用 E-R 图描述，如图3-5所示：一个用户可提交多个跟踪任务，用户与任务"
             "之间为 1:N 联系。")
    add_fig(doc, F("fig4_1_数据库E-R图.png"), "图3-5 数据库 E-R 图", width_cm=13.5)
    add_body(doc,
             "数据库物理结构设计采用 MySQL 语法规范设计，库名 zongshe3_track，字符集 utf8mb4。为便于"
             "跨电脑演示与测试，系统默认以内嵌 H2 数据库运行（H2 以 MySQL 兼容模式启动，免安装、"
             "零外部依赖，数据库文件自动生成于运行目录），同时保留 MySQL 配置（--spring.profiles"
             ".active=mysql 可切换，需先执行 sql/init.sql）。各数据表结构如表3-3、表3-4所示（表中 "
             "PK 为主键、FK 为外键、Not null 为非空）。")
    add_table(doc, "表3-3 用户表 t_user 结构",
              ["字段名", "中文名", "类型", "主键/外键", "备注"],
              [
                  ["id", "用户ID", "BIGINT", "PK", "自增主键"],
                  ["username", "用户名", "VARCHAR(50)", "", "唯一 uk_username, Not null"],
                  ["password", "密码", "VARCHAR(100)", "", "BCrypt 加密, Not null"],
                  ["nickname", "昵称", "VARCHAR(50)", "", "可空"],
                  ["role", "角色", "VARCHAR(10)", "", "USER/ADMIN, Not null, 默认 USER"],
                  ["create_time", "注册时间", "DATETIME", "", "默认当前时间"],
              ],
              col_widths=[2.8, 2.2, 3.2, 2.2, 3.6])
    add_table(doc, "表3-4 任务表 t_task 结构",
              ["字段名", "中文名", "类型", "主键/外键", "备注"],
              [
                  ["id", "任务ID", "BIGINT", "PK", "自增主键"],
                  ["user_id", "提交用户ID", "BIGINT", "FK→t_user.id", "Not null, 索引 idx_user"],
                  ["media_type", "媒体类型", "VARCHAR(10)", "", "VIDEO/IMAGE"],
                  ["file_name", "原始文件名", "VARCHAR(255)", "", "—"],
                  ["file_path", "原文件路径", "VARCHAR(500)", "", "服务器绝对路径"],
                  ["bbox", "首帧目标框", "VARCHAR(50)", "", "x,y,w,h；空=自动"],
                  ["status", "任务状态", "VARCHAR(20)", "", "PENDING/PROCESSING/SUCCESS/FAILED"],
                  ["result_path", "结果路径", "VARCHAR(500)", "", "Web 访问路径 /files/..."],
                  ["log_path", "日志路径", "VARCHAR(500)", "", "跟踪轨迹日志 JSON"],
                  ["stats_json", "统计JSON", "TEXT", "", "帧数/丢失/恢复/FPS 等"],
                  ["error_msg", "错误信息", "VARCHAR(500)", "", "失败原因"],
                  ["create_time", "提交时间", "DATETIME", "", "默认当前时间"],
                  ["finish_time", "完成时间", "DATETIME", "", "处理结束时间"],
              ],
              col_widths=[2.8, 2.6, 2.6, 2.6, 3.6])

    add_h3(doc, "3.2.4 系统接口设计")
    add_body(doc,
             "系统对外 REST 接口采用统一返回体 {code, msg, data}，code=0 表示成功；业务接口请求头"
             "携带 Authorization: Bearer <JWT>。主要接口如表3-5所示。")
    add_table(doc, "表3-5 系统主要 REST 接口清单",
              ["方法", "接口路径", "功能", "权限"],
              [
                  ["POST", "/api/auth/register", "用户注册", "公开"],
                  ["POST", "/api/auth/login", "用户登录，返回 JWT", "公开"],
                  ["GET", "/api/user/info", "查询个人信息", "登录用户"],
                  ["PUT", "/api/user/password", "修改密码", "登录用户"],
                  ["POST", "/api/task/upload", "上传视频/图片并创建任务(multipart)；defer=true 时只"
                   "提取首帧与自动目标框（两段式第一步）", "登录用户"],
                  ["POST", "/api/task/{id}/start", "两段式第二步：带框选结果开始处理", "登录用户(限本人)"],
                  ["GET", "/api/task/{id}", "查询任务状态与结果（含归属校验）", "登录用户(限本人)"],
                  ["GET", "/api/task/list", "我的任务记录(分页)", "登录用户"],
                  ["GET", "/api/admin/users", "用户分页查询", "ADMIN"],
                  ["POST", "/api/admin/user", "新增用户", "ADMIN"],
                  ["PUT", "/api/admin/user/{id}/reset", "重置密码", "ADMIN"],
                  ["DELETE", "/api/admin/user/{id}", "删除用户(级联删除任务)", "ADMIN"],
                  ["GET", "/api/admin/tasks", "任务记录分页/状态过滤", "ADMIN"],
                  ["DELETE", "/api/admin/task/{id}", "删除任务记录", "ADMIN"],
                  ["GET", "/api/admin/stats", "系统统计卡片", "ADMIN"],
                  ["GET", "/files/**", "访问结果视频/结果图片/日志文件", "公开(静态映射)"],
              ],
              col_widths=[2.2, 5.2, 4.4, 2.6])
    add_body(doc,
             "视觉服务（Python + FastAPI）独立提供以下接口，由后端 TaskService 通过 Java HttpClient "
             "调用：GET /api/v1/health（健康检查）、POST /api/v1/prepare（提取首帧图片并返回素材原始"
             "尺寸与自动识别目标框，供小程序框选流程使用）、POST /api/v1/track_local（输入本地素材"
             "路径与 bbox，返回结果文件路径与统计信息）。")

    # ---------------- 3.3 系统实现 ----------------
    add_h2(doc, "3.3 系统实现")
    add_h3(doc, "3.3.1 视觉跟踪模块实现")
    add_body(doc,
             "视觉跟踪模块核心类 SingleObjectTracker 维护状态机并封装参考模型（灰度模板 + HSV 直方图）。"
             "遮挡检测与状态切换的核心代码如下（节选，完整源码见 01-代码/vision/tracker.py）：")
    add_code(doc, '''# 跟踪质量校验：周期性直方图相关度 + 逐帧越界检查，判定是否进入搜索
if self.stats["total_frames"] % VERIFY_EVERY == 0:
    cx0, cy0 = max(0, x), max(0, y)
    cx1 = min(w, x + bw); cy1 = min(h, y + bh)
    if cx1 - cx0 > 4 and cy1 - cy0 > 4:
        sim = self._hist_similarity(frame[cy0:cy1, cx0:cx1])
    else:
        sim = 0.0
    # 只有校验帧才更新计数：否则会被中间帧的默认相似度冲掉，LOST_CONFIRM 永远达不到
    self._lost_counter = self._lost_counter + 1 if sim < HIST_LOST_TH else 0
self._oob_counter = self._oob_counter + 1 if out_of_bounds else 0

if self._lost_counter >= LOST_CONFIRM or self._oob_counter >= OOB_CONFIRM:
    self._enter_search()                       # 连续3次校验/3帧越界确认 -> 搜索
    return self._search(frame)
# 未达确认阈值：继续输出跟踪框（防抖，避免单帧抖动就整帧全图搜索）
self._last_bbox = (x, y, bw, bh)''', caption="代码3-1 遮挡检测与状态切换（tracker.py 节选）")
    add_body(doc,
             "搜索模式采用“多尺度模板匹配 + 直方图验证 + 连续帧确认”三重复检测，命中后重新初始化"
             "跟踪器，实现遮挡后的再捕捉：")
    add_code(doc, '''def _search(self, frame):
    self.stats["searching_frames"] += 1
    best = self._detect(frame)                 # 多尺度模板匹配
    if best is None:
        self._search_hits = 0
        return None, self.state
    sim = self._hist_similarity(frame[best[1]:best[1]+best[3],
                                      best[0]:best[0]+best[2]])
    if sim >= HIST_FOUND_TH:                   # 直方图二次验证
        self._search_hits += 1
        if self._search_hits >= SEARCH_CONFIRM:    # 连续2帧确认找回
            self._new_tracker(frame, best)     # 重新初始化跟踪器
            self.state = TrackerState.TRACKING
            self.stats["recoveries"] += 1      # 累计恢复次数
            return best, self.state
    else:
        self._search_hits = 0
    return best, self.state''', caption="代码3-2 遮挡搜索与自动恢复（tracker.py 节选）")

    add_h3(doc, "3.3.2 Web 后端与管理端实现")
    add_body(doc,
             "Web 后端采用 Spring Boot 3 + MyBatis 注解式 SQL。任务上传接口将文件落盘后立即登记任务"
             "并返回任务 ID，视觉处理在单线程线程池中异步执行，核心代码如下（节选，完整源码见"
             "01-代码/server）：")
    add_code(doc, '''// 上传：落盘 + 登记任务(PENDING)；defer=true 时只提首帧供前端画框
@PostMapping("/upload")
public Result<TrackTask> upload(@RequestAttribute("uid") Long uid,
        @RequestParam("file") MultipartFile file,
        @RequestParam(value = "bbox", required = false) String bbox,
        @RequestParam(value = "defer", required = false, defaultValue = "false") boolean defer) {
    return Result.ok(taskService.upload(uid, file, bbox, defer));
}
// 两段式第二步：带框选结果开始处理（含任务归属校验）
@PostMapping("/{id}/start")
public Result<TrackTask> start(@PathVariable Long id,
        @RequestAttribute("uid") Long uid,
        @RequestAttribute("role") String role,
        @RequestBody(required = false) Map<String, String> body) {
    return Result.ok(taskService.start(id, body == null ? null : body.get("bbox"),
                                       uid, role));
}
// TaskService.process(): PROCESSING -> 调视觉服务 -> SUCCESS/FAILED
private void process(Long taskId) {
    taskMapper.markProcessing(taskId);
    try {
        Map<String, String> r = visionClient.track(task.getFilePath(), task.getBbox());
        // 图片素材输出 *_tracked.jpg，视频素材输出 *_tracked.mp4
        String name = taskId + ("IMAGE".equals(task.getMediaType())
                                ? "_tracked.jpg" : "_tracked.mp4");
        Files.copy(Paths.get(r.get("resultVideo")), vidDst, REPLACE_EXISTING);
        Files.copy(Paths.get(r.get("logFile")),     logDst, REPLACE_EXISTING);
        taskMapper.markSuccess(taskId, webUrlOf(vidDst), webUrlOf(logDst),
                               r.get("statsJson"));
    } catch (Exception e) {
        taskMapper.markFailed(taskId, e.getMessage());   // 失败留痕
    }
}''', caption="代码3-3 上传、两段式框选与异步任务调度（TaskController/TaskService 节选）")
    add_body(doc,
             "管理端 Web 页面以原生单页实现（admin.html + admin.js），零构建依赖，随后端静态资源发布；"
             "采用 JWT 拦截器统一鉴权（/api/admin/** 要求 ADMIN 角色），任务详情与任务启动接口在业务"
             "层再做一次归属校验（普通用户只能访问本人任务），页面运行效果见 3.3.4 节。")

    add_h3(doc, "3.3.3 小程序端实现")
    add_body(doc,
             "小程序端共包含登录、注册、上传跟踪、处理结果、我的记录、个人中心 6 个页面。网络层统一"
             "封装在 utils/request.js 中：wx.request 自动携带 JWT，401 时自动跳转登录；上传通过 "
             "wx.uploadFile 携带 Authorization 头与 bbox 表单字段。"
             "由于小程序无法把 <video> 组件的画面绘制到 canvas，做不到像 Web 端那样在前端取帧，"
             "因此采用“后端提首帧 + 前端在图片上画框”的方案：选好素材后以 defer=true 上传，"
             "后端返回首帧图片 URL、素材原始尺寸与自动识别目标框；前端在 <image> 上叠加同层 <view>，"
             "用 touchstart/move/end 实现手指拖拽画框，并把显示坐标按“原始尺寸/显示尺寸”的比例换算"
             "回原始像素。上传与框选的核心代码如下（节选）：")
    add_code(doc, '''// 选素材后上传并提取首帧（index.js）
const task = await upload('/api/task/upload?defer=true', videoPath, {});
// task.frameUrl / task.imgWidth / task.imgHeight / task.autoBbox 用于显示与画框
// 手指拖拽画框：显示坐标 -> 素材原始像素
onTouchEnd() {
    const s = this.data.imgW / this.data.dispW;     // 原始像素 / 显示像素
    const x = Math.round(p.x * s), y = Math.round(p.y * s);
    const w = Math.round(p.w * s), h = Math.round(p.h * s);
    this.setBox({ x, y, w, h }, false);             // 裁剪并保证 >= 5px
}
// 提交：带框选结果开始处理（index.js）
async submit() {
    const box = this._box;
    const bbox = box ? (box.x + ',' + box.y + ',' + box.w + ',' + box.h) : '';
    await request('POST', '/api/task/' + taskId + '/start', { bbox: bbox });
    wx.navigateTo({ url: '/pages/result/result?id=' + taskId });
}
// 结果页轮询（result.js）：处理中每1.5s查询一次，完成后展示
const t = await request('GET', '/api/task/' + this.data.taskId);
if (t.status === 'PENDING' || t.status === 'PROCESSING') {
    setTimeout(() => this.poll(), 1500);   // 继续轮询
} else {
    this.setData({ status: t.status, resultUrl: getUrl(t.resultPath) });
}''', caption="代码3-4 小程序两段式上传、首帧框选与状态轮询（节选）")

    add_h3(doc, "3.3.4 运行界面与系统测试截图")
    add_body(doc,
             "视觉跟踪模块对合成测试视频（640×480、30 帧/秒、420 帧，画面中央存在遮挡墙）的运行"
             "结果如图3-6~图3-10所示。测试视频中目标（蓝色圆球）自左向右运动，途中被灰色遮挡墙"
             "完全遮挡约 0.9 秒，系统先进入 SEARCHING 搜索状态，目标重新出现后在第 262 帧成功"
             "找回并继续跟踪直至视频结束，完整展示了“遮挡丢失—自动重检测—恢复跟踪”的能力。")
    add_fig(doc, S("shot1_源视频首帧_目标待框选.png"), "图3-6 测试视频首帧（目标为蓝色圆球）", width_cm=10.5)
    add_fig(doc, S("shot2_跟踪中_简单运动.png"), "图3-7 正常跟踪状态（绿色跟踪框与运动轨迹）", width_cm=10.5)
    add_fig(doc, S("shot3_目标被遮挡_进入搜索.png"), "图3-8 目标被遮挡丢失，进入搜索状态", width_cm=10.5)
    add_fig(doc, S("shot4_遮挡后重新捕获_继续跟踪.png"), "图3-9 目标重现后被自动重新捕获并继续跟踪", width_cm=10.5)
    add_fig(doc, S("shot5_跟踪至末尾_轨迹完整.png"), "图3-10 跟踪至视频末尾，运动轨迹完整", width_cm=10.5)

    add_body(doc,
             "Web 管理端运行界面如图3-11~图3-14所示，管理员登录后可查看统计卡片、进行用户管理"
             "（查询/新增/删除/重置密码）、查看全部任务记录并在线播放跟踪结果视频。")
    add_fig(doc, S("web_01_登录页.png"), "图3-11 Web 管理端登录页", width_cm=12.5)
    add_fig(doc, S("web_02_用户管理.png"), "图3-12 Web 管理端-用户管理（统计卡片+用户列表）", width_cm=13.5)
    add_fig(doc, S("web_03_任务记录.png"), "图3-13 Web 管理端-任务记录", width_cm=13.5)
    add_fig(doc, S("web_04_结果弹窗.png"), "图3-14 在线播放跟踪结果视频与统计", width_cm=13.5)

    add_body(doc,
             "小程序端各页面结构如图3-15~图3-17所示（界面原型示意，真实运行截图请按 README 中的"
             "截图清单在微信开发者工具中运行小程序后替换）：")
    add_fig(doc, F("mini_1_登录页原型.png"), "图3-15 小程序-登录页", width_cm=5.2)
    add_fig(doc, F("mini_2_上传页原型.png"), "图3-16 小程序-上传跟踪页", width_cm=5.2)
    add_fig(doc, F("mini_3_结果页原型.png"), "图3-17 小程序-处理结果页", width_cm=5.2)

    add_body(doc,
             "为便于答辩与日常演示，系统还提供了网页检测端（/track.html），作为小程序功能的等价"
             "浏览器入口：支持将视频或图片文件【拖拽】到页面、点击选择或直接 Ctrl+V 粘贴；可选择"
             "“提取视频首帧并在画布上拖拽框选目标”（框选坐标自动换算为素材原始像素）或使用自动"
             "识别目标（视频=运动目标检测，图片=画面中央）；图片素材本身即首帧，选好后直接进入"
             "画框步骤。提交后页面自动轮询任务状态，处理完成后直接展示结果视频（或结果图片）与"
             "跟踪统计（帧数/速度/丢失/找回等），并支持本人历史记录回看。该页面无需安装任何客户端，"
             "打开浏览器即可完成“提交素材—跟踪检测—结果回看”的完整闭环，运行界面如图3-18、"
             "图3-19所示。")
    add_fig(doc, S("web_05_网页检测端_拖拽上传.png"), "图3-18 网页检测端-拖拽上传与目标框选", width_cm=13.5)
    add_fig(doc, S("web_06_网页检测端_检测结果.png"), "图3-19 网页检测端-检测结果展示", width_cm=13.5)


# ============================ 第4章 ============================
def build_ch4(doc):
    add_h1(doc, "4 系统测试")
    add_h2(doc, "4.1 测试环境")
    add_table(doc, "表4-1 系统测试环境",
              ["项目", "配置"],
              [
                  ["操作系统", "Windows 11（64 位）"],
                  ["Web 后端", "Spring Boot 3.0.2 / JDK 21（编译目标 Java 17）/ MyBatis"],
                  ["数据库", "内嵌 H2（MySQL 兼容模式，默认，免安装）；亦可切换 MySQL 8.0"],
                  ["视觉处理", "Python 3.12 + OpenCV 4.11（contrib，含 CSRT）+ FastAPI + Uvicorn"],
                  ["测试素材", "合成视频 A：640×480、30fps、420 帧、无遮挡（场景A）；"
                   "合成视频 B：同规格、画面中央含遮挡墙（场景B）；测试图片 640×480（图片链路）"],
                  ["浏览器/工具", "Edge（Web 管理端）、微信开发者工具（小程序）、"
                   "自研接口回归脚本（50 项断言）/命令行接口测试"],
              ],
              col_widths=[3.2, 10.8])
    add_h2(doc, "4.2 测试记录")
    add_h3(doc, "4.2.1 功能测试")
    add_body(doc,
             "按功能模块设计测试用例并逐项执行，结果如表4-2所示（参照模板表6.1格式）。")
    add_table(doc, "表4-2 功能测试用例及结果",
              ["功能模块", "操作描述", "预期结果", "实际结果"],
              [
                  ["用户注册", "输入重复用户名注册", "提示“用户名已被注册”", "通过"],
                  ["用户注册", "密码长度不足6位注册", "提示密码长度不足", "通过"],
                  ["用户注册", "合法信息注册", "注册成功并可登录", "通过"],
                  ["用户登录", "错误密码登录", "提示“用户名或密码错误”", "通过"],
                  ["用户登录", "正确账号登录", "返回 JWT 与用户信息", "通过"],
                  ["修改密码", "原密码错误提交", "提示原密码错误", "通过"],
                  ["修改密码", "正确修改后重新登录", "新密码可登录", "通过"],
                  ["素材上传", "上传非视频/图片格式文件", "提示仅支持 mp4 等视频与 jpg 等图片",
                   "通过"],
                  ["素材上传", "bbox 格式非法", "提示 bbox 应为 x,y,w,h", "通过"],
                  ["素材上传", "正常上传视频+自动识别目标", "返回任务ID，状态轮询至 SUCCESS，"
                   "结果视频可播放", "通过"],
                  ["素材上传", "上传图片素材（两段式框选）", "返回首帧与自动框，start 后输出"
                   "标注结果图片", "通过"],
                  ["结果查询", "任务完成后拉取任务详情", "返回结果文件路径与统计 JSON", "通过"],
                  ["结果播放", "访问 /files/results/*.mp4", "视频可正常播放", "通过"],
                  ["结果查看", "访问 /files/results/*_tracked.jpg", "结果图片可正常显示", "通过"],
                  ["用户管理", "管理员新增/查询用户", "列表实时更新", "通过"],
                  ["用户管理", "重置密码后登录", "新密码可登录", "通过"],
                  ["用户管理", "删除用户", "用户与任务记录级联删除", "通过"],
                  ["权限控制", "普通用户访问管理端接口", "返回 403 无权限", "通过"],
                  ["权限控制", "用户 B 查询/启动用户 A 的任务", "返回 403 无权访问该任务", "通过"],
                  ["任务管理", "状态过滤查询任务", "仅返回对应状态记录", "通过"],
                  ["统计卡片", "查看系统统计", "数量与数据库一致", "通过"],
              ],
              col_widths=[2.6, 5.2, 4.2, 2.0])

    add_h3(doc, "4.2.2 视觉跟踪专项测试")
    add_body(doc,
             "针对任务书的核心要求“目标被短暂遮挡后再次出现能再次捕捉继续跟踪”，使用"
             "make_demo_video.py 生成两段对照视频（场景A 不含遮挡墙、场景B 含遮挡墙，其余完全一致），"
             "设计三种典型场景对视觉模块进行专项测试，结果如表4-3所示（运行日志见 "
             "demo/vision_out/*_log.json）。")
    add_table(doc, "表4-3 视觉跟踪专项测试记录",
              ["测试场景", "场景描述", "预期结果", "实际结果"],
              [
                  ["场景A：简单运动跟踪", "demo_ball_seamless.mp4：目标无遮挡匀速运动 420 帧",
                   "全程 TRACKING，无丢失", "通过：tracking 420 帧、搜索 0 帧、丢失 0 次，"
                   "最终 TRACKING"],
                  ["场景B：短暂完全遮挡后恢复", "demo_ball_track.mp4：目标运动中被遮挡墙完全"
                   "遮挡约 0.9s 后重现",
                   "丢失后自动重检测并找回，继续跟踪至结束",
                   "通过：tracking 356 帧、搜索 64 帧、丢失 1 次、恢复 1 次，最终 TRACKING"],
                  ["场景C：目标轨迹跟踪完整性", "目标自画面左端运动至右端（第55帧 x≈55 → "
                   "第419帧 x≈585）",
                   "轨迹连续、末帧跟踪框贴合目标",
                   "通过：恢复后跟踪框回到目标上（第 262 帧 bbox=362,192,52,54），"
                   "末帧仍为 TRACKING"],
                  ["性能：处理速度", "640×480@30fps，420 帧",
                   "实时（≥25 帧/秒）",
                   "通过：场景A 约 104 帧/秒；场景B（含 64 帧全图搜索）约 65 帧/秒"],
                  ["图片素材链路", "640×480 测试图片，指定 bbox=(150,120,80,80)",
                   "输出标注了目标框的结果图片并给出统计",
                   "通过：result=*_tracked.jpg，stats.media_type=IMAGE，"
                   "target_bbox=150,120,80,80"],
              ],
              col_widths=[3.4, 4.4, 3.4, 4.0])
    add_body(doc,
             "专项测试中场景B的完整过程数据：视频共 420 帧，正常跟踪 356 帧，搜索 64 帧，丢失事件"
             "1 次，自动重新捕获 1 次，最终状态为 TRACKING。目标自约第 200 帧起被遮挡墙完全遮挡，"
             "系统于第 199 帧检测到丢失进入搜索状态，第 262 帧目标重新出现后被成功找回并恢复跟踪"
             "（bbox 由搜索得到的 362,192,52,54 重新初始化跟踪器），验证了直方图质量校验与模板匹配"
             "重检测机制的有效性。")

    add_h2(doc, "4.3 测试结论及分析")
    add_body(doc,
             "功能测试 21 项用例全部通过，覆盖注册、登录、改密、视频与图片上传、两段式框选、轮询、"
             "结果播放/查看、用户管理、任务管理、权限与越权防护等全部功能点；视觉专项测试证明系统"
             "在简单运动与遮挡恢复两类场景下均能稳定跟踪，遮挡恢复机制有效；性能上处理速度 65~104 "
             "帧/秒，满足实时性要求。")
    add_body(doc,
             "质量评定：系统功能完善、界面友好、权限控制到位，视觉核心能力满足任务书要求，代码"
             "命名规范、注释完整，测试过程留痕充分，可进行答辩演示。主要局限在于：CSRT 对目标形变"
             "较大或长时间消失场景鲁棒性有限；重检测依赖初始参考模板，目标外观剧烈变化时可能找回"
             "失败，后续可引入在线模板更新与深度学习跟踪器改进。")

    add_h2(doc, "4.4 测试结论")
    add_body(doc,
             "综合以上功能测试、视觉专项测试与性能测试结果，系统各项指标均达到设计要求，可以交付。"
             "对测试中暴露出的问题已修复并在回归测试中验证通过，主要包括：① 完全遮挡场景下跟踪框"
             "漂移导致恢复较慢——通过调整判定阈值与连续帧确认机制解决；② 丢失确认的防抖计数被非"
             "校验帧重置、LOST_CONFIRM 形同虚设——改为仅在质量校验帧更新计数后，同一测试视频的搜索"
             "帧数由 74 帧降到 64 帧、正常跟踪帧数由 346 帧提升到 356 帧；③ 任务详情与任务启动接口"
             "缺少归属校验（可越权访问他人任务）——补充归属校验后返回 403；④ 图片素材此前会被当作"
             "视频处理而失败——新增图片处理链路，输出标注目标框的结果图片。")


# ============================ 第5章 ============================
def build_ch5(doc):
    add_h1(doc, "5 系统总结")
    add_h2(doc, "5.1 系统特色")
    add_body(doc,
             "（1）三端联动、架构解耦：系统完整实现了“小程序用户端 + Web 管理端 + Python 视觉"
             "服务端 + 数据库/文件存储”的任务书要求架构，视觉处理层独立成服务，通过 REST 接口与"
             "业务后端解耦，各端职责清晰、可独立部署。")
    add_body(doc,
             "（2）遮挡自动重检测是本系统最大特色：针对普通跟踪器“目标一挡就丢”的痛点，设计了"
             "“跟踪质量校验 + 搜索状态机 + 多尺度模板匹配 + 直方图验证 + 连续帧确认”的恢复机制，"
             "实测目标被完全遮挡约 0.9 秒后能自动找回并继续跟踪，直接对应任务书要求。")
    add_body(doc,
             "（3）素材类型完整、跨端一致：任务书要求“提交要处理的图片或者视频”，系统对两种素材"
             "统一了处理链路——视频逐帧跟踪并输出带轨迹的结果视频，图片在首帧上定位目标并输出标注"
             "结果图；两端（小程序 / 网页检测端）都能在首帧画面上拖拽框选目标，框选坐标统一换算回"
             "素材原始像素，保证“所见即所框”。")
    add_body(doc,
             "（4）工程化完整：遵循软件工程规范完成需求分析、设计、编码、测试全过程；代码命名规范、"
             "文件头注明作者学号与功能说明；使用 Git 团队协作；接口统一、错误留痕、权限分级（JWT + "
             "ADMIN 角色校验 + 任务归属校验），可作为规范的中小型软件项目范例。")
    add_body(doc,
             "（5）轻量可复现：视觉模块免深度学习训练、离线可运行；数据库默认内嵌 H2 免安装；"
             "管理端零构建依赖；后端随仓库附带可执行 jar 并提供免 Maven 重新打包脚本；配合合成测试"
             "视频与图片，可在任意一台装有 JDK 17+ 与 Python 3.10+ 的电脑上快速复现演示。")

    add_h2(doc, "5.2 展望工作")
    add_body(doc,
             "（1）算法升级：将 CSRT 替换为基于深度学习的 SiamRPN/ByteTrack 等跟踪器，并引入在线"
             "模板更新与多特征融合，进一步提升形变、长时间遮挡场景下的鲁棒性；")
    add_body(doc,
             "（2）实时视频流：当前为“上传素材—离线处理”模式，后续可接入 RTSP 摄像头推流与"
             "WebSocket 进度推送，实现真正意义的实时跟踪与云端转发；")
    add_body(doc,
             "（3）交互体验：在已实现的首帧拖拽框选基础上，增加框的边角微调手柄、多目标候选列表、"
             "处理进度条与任务删除、结果下载（含结果图片一键保存）等能力；")
    add_body(doc,
             "（4）系统完善：引入用户注册邮箱/手机校验、操作审计日志、消息推送，并完成容器化"
             "（Docker Compose）一键部署，使系统更接近生产级应用。")

    add_h2(doc, "5.3 使用说明")
    add_body(doc,
             "系统部署与启动步骤如下（详细说明见工程根目录 README.md）。"
             "推荐方式：直接双击工程根目录的 启动-网页检测端.bat（或 启动-管理端.bat），"
             "脚本会自动探测 Python/JDK、检查并安装 opencv-contrib-python、按需启动视觉服务与 Web "
             "后端，最后打开浏览器，无需手工敲命令；数据库默认使用内嵌 H2，无需安装 MySQL。")
    add_table(doc, "表5-1 系统启动步骤",
              ["步骤", "操作", "说明"],
              [
                  ["1", "一键启动（推荐）", "双击 启动-网页检测端.bat 或 启动-管理端.bat；脚本自动"
                   "完成环境探测、依赖安装、服务启动与打开浏览器"],
                  ["2", "初始化数据库（可选）", "默认使用内嵌 H2，随程序自动建库建表，无需任何操作；"
                   "若改用 MySQL，先执行 server/sql/init.sql 创建 zongshe3_track 库与数据表，"
                   "再以 --spring.profiles.active=mysql 启动"],
                  ["3", "启动视觉服务", "python api_service.py（默认 127.0.0.1:9000）；"
                   "依赖：pip install -r requirements.txt（含 opencv-contrib-python，CSRT 必需）"],
                  ["4", "启动 Web 后端", "java -jar target/track-server-1.0.0.jar（默认 8080，"
                   "首次启动自动创建管理员 admin/123456）"],
                  ["5", "打开管理端", "浏览器访问 http://localhost:8080/admin.html，用 admin 登录"],
                  ["6", "运行小程序", "微信开发者工具导入 01-代码/miniprogram，勾选“不校验合法域名”，"
                   "登录后选择视频或图片，在首帧上拖拽框选目标后提交"],
                  ["7", "本地视觉演示", "python gui_track.py --video ../../demo/demo_ball_track.mp4"
                   "（鼠标框选目标后回车开始跟踪，也支持 --camera 0 摄像头跟踪）"],
              ],
              col_widths=[1.4, 3.6, 9.0])
    add_body(doc,
             "测试数据与演示素材位于 demo/ 目录：demo_ball_track.mp4 为含遮挡场景的合成测试视频"
             "（场景B），demo_ball_seamless.mp4 为同规格无遮挡对照视频（场景A），"
             "demo/vision_out/ 保存处理结果与逐帧日志，demo/screenshots/ 与 demo/figures/ 保存报告"
             "与答辩所需截图和示意图。")


# ============================ 第6章 ============================
def build_ch6(doc):
    add_h1(doc, "6 心得体会")
    add_body(doc,
             "通过本次综合实践，我完整经历了一个“三端联动 + 机器视觉”复杂工程问题的分析、设计、"
             "实现、测试与总结全流程，收获颇多。")
    add_body(doc,
             "首先是“复杂工程问题”的拆解能力得到锻炼。课程设计开始时面对任务书中小程序、Web、"
             "机器视觉三座“大山”，我们首先把系统拆分为相对独立的模块（视觉、后端、小程序、管理端），"
             "约定好 REST 接口与数据表结构后再并行开发，这让我切实体会到接口契约与模块化设计在团队"
             "项目中的价值。")
    add_body(doc,
             "其次是视觉算法的深入理解。CSRT 相关滤波跟踪器本身并不复杂，但真正把“遮挡后重新捕捉”"
             "做成稳定可用却需要反复打磨：最初测试视频的遮挡墙比目标还窄，目标从未被完全遮挡，"
             "导致丢失检测根本不触发；将遮挡场景修正后，又面临阈值灵敏度与误判的平衡。最终通过"
             "“直方图质量校验 + 多尺度模板匹配 + 连续帧确认”的方案解决问题，这个过程让我深刻理解了"
             "算法参数设计与实验验证的重要性。")
    add_body(doc,
             "再次是工程规范与团队协作的实践。项目全程使用 Git 进行版本控制与分工协作，代码按规范"
             "命名并书写注释，每个源文件头注明作者学号与功能；联调阶段我们按“视觉服务—后端—前端”"
             "顺序逐层打通，任何接口异常都能快速定位到具体模块。此外，测试记录与文档同步撰写保证了"
             "“代码—测试—文档”的一致，答辩时也能清晰说明每一个模块的原理。")
    add_body(doc,
             "最后，我也认识到自身不足：对深度学习跟踪算法的理解还不够深入，服务端并发与安全加固"
             "（如文件类型白名单、限流）仍有提升空间。本次实践为后续生产实习与毕业设计打下了坚实"
             "基础，我将继续保持严谨求实的态度，把课程中学到的工程化方法运用到今后的学习与工作中。")


# ============================ 参考文献 ============================
def build_refs(doc):
    add_h1(doc, "参考文献")
    refs = [
        "[1] Henriques J F, Caseiro R, Martins P, et al. High-Speed Tracking with "
        "Kernelized Correlation Filters[J]. IEEE Transactions on Pattern Analysis and "
        "Machine Intelligence, 2015, 37(3): 583-596.",
        "[2] Lukezic A, Vojir T, Cehovin Zajc L, et al. Discriminative Correlation Filter "
        "Tracker with Channel and Spatial Reliability[J]. International Journal of "
        "Computer Vision, 2018, 126(7): 671-688.",
        "[3] Bradski G. The OpenCV Library[J]. Dr. Dobb's Journal of Software Tools, "
        "2000, 25(11): 120-125.",
        "[4] 毛星云, 冷雪飞, 王碧辉, 等. OpenCV3编程入门[M]. 北京: 电子工业出版社, 2015.",
        "[5] 王珊, 萨师煊. 数据库系统概论(第5版)[M]. 北京: 高等教育出版社, 2014.",
        "[6] OpenCV.org. OpenCV Tutorials: Video Analysis — Object Tracking[EB/OL]. "
        "https://docs.opencv.org/4.x/d9/df8/tutorial_root.html.",
        "[7] 微信开放文档. 微信小程序开发文档[EB/OL]. "
        "https://developers.weixin.qq.com/miniprogram/dev/framework/.",
        "[8] Spring 官方文档. Spring Boot Reference Documentation[EB/OL]. "
        "https://docs.spring.io/spring-boot/docs/current/reference/html/.",
    ]
    for r in refs:
        add_body(doc, r, indent=0)


# ============================ 附录 ============================
def build_appendix(doc):
    add_h1(doc, "附录A 项目源代码目录结构")
    add_code(doc, '''综设3-单目标跟踪/
├── 01-代码/
│   ├── vision/                     # 视觉处理模块（Python + OpenCV）
│   │   ├── tracker.py              # 跟踪器核心：CSRT + 遮挡检测/重检测状态机 + 图片定位
│   │   ├── gui_track.py            # 本地交互演示：框选目标/摄像头跟踪
│   │   ├── api_service.py          # FastAPI 视觉服务（prepare/track_local/track/health）
│   │   ├── make_demo_video.py      # 合成测试视频（--no_wall 生成无遮挡对照视频）
│   │   ├── requirements.txt        # Python 依赖（含 opencv-contrib-python，CSRT 必需）
│   │   └── extract_screens.py      # 从结果视频抽取报告截图
│   ├── server/                     # Web 后端（Spring Boot 3 + MyBatis）
│   │   ├── pom.xml
│   │   ├── sql/init.sql            # MySQL 建库建表脚本
│   │   ├── target/track-server-1.0.0.jar   # 已构建的可执行 jar（双击即用）
│   │   ├── src/main/java/com/zongshe3/track/
│   │   │   ├── controller/         # Auth/User/Task/Admin/Root 控制器
│   │   │   ├── service/            # 业务逻辑 + VisionClient(调视觉服务)
│   │   │   ├── mapper/             # MyBatis 注解式 Mapper
│   │   │   ├── interceptor/        # JWT 认证与角色鉴权
│   │   │   ├── config/             # Web 配置/静态映射/管理员初始化
│   │   │   ├── common/             # Result/JwtUtil/BizException/全局异常
│   │   │   └── pojo/               # 实体类 User/TrackTask
│   │   ├── src/main/resources/
│   │   │   ├── application*.yml    # 主配置 + h2/mysql 两套 profile
│   │   │   ├── db/schema-h2.sql    # 内嵌 H2 建表脚本
│   │   │   └── static/             # Web 页面（admin.html/track.html/css/js）
│   │   └── uploads/                # 运行期上传原文件/首帧/结果（自动生成）
│   └── miniprogram/                # 微信小程序端
│       ├── app.js/app.json/app.wxss
│       ├── utils/request.js        # 请求/上传封装（携带JWT）
│       └── pages/                  # login/register/index/result/records/mine
├── 02-文档/                        # 课程设计报告 docx 与预览 PDF
├── 03-答辩PPT/                     # 答辩 PPT
├── demo/                           # 测试视频/结果/截图/示意图
│   ├── demo_ball_track.mp4         # 场景B：含遮挡墙
│   ├── demo_ball_seamless.mp4      # 场景A：无遮挡对照
│   ├── vision_out/                 # 跟踪结果与逐帧日志
│   ├── screenshots/                # 报告与答辩用真实截图
│   └── figures/                    # 示意图（架构/用例/状态机等）
├── tools/                          # 启动/打包/图/报告/PPT 生成脚本
│   ├── start_all.ps1               # 一键启动（环境探测+依赖安装+服务启动）
│   ├── build_jar_nomaven.ps1       # 无 Maven 环境下重新打包 jar
│   ├── gen_report.py、report_part1/2.py、docx_lib.py   # 报告生成
│   └── gen_ppt.py、figures.py                          # PPT 与示意图生成
├── 启动-网页检测端.bat / 启动-管理端.bat   # 双击启动并打开浏览器
└── README.md                       # 部署与使用说明''')
